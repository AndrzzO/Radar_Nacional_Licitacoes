"""Validação opcional REAL; não usa fixtures nem altera o banco do usuário."""
import os,sys,json,time
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]; sys.path.insert(0,str(ROOT))
DEST=ROOT/'scratch'/'validacao_online'; DEST.mkdir(parents=True,exist_ok=True)
os.environ['RADAR_DATA_DIR']=str(DEST); os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
from PySide6.QtWidgets import QApplication,QMessageBox
from src.database.connection import DatabaseConnection
from src.database.migrations import init_database
from src.database.repositories import ProcurementRepository,LocationRepository
from src.application.location_service import FALLBACK_ESTADOS
from src.application.export_service import ExportService
from src.domain.filters import LicitacaoFilter
from src.ui.main_window import MainWindow
from src.utils.datetime_utils import get_current_time_sp
from openpyxl import load_workbook

app=QApplication([])
from PySide6.QtGui import QFontDatabase,QFont
QFontDatabase.addApplicationFont('C:/Windows/Fonts/segoeui.ttf'); app.setFont(QFont('Segoe UI',9))
db=DatabaseConnection(DEST/'online.sqlite3')
with db.session() as c: init_database(c)
LocationRepository(db).save_estados(FALLBACK_ESTADOS)
window=MainWindow(db); window.show(); window._navigate(1)
def wait(worker):
 while worker and worker.isRunning(): app.processEvents(); time.sleep(.02)
 app.processEvents()
wait(window.page_search.query_worker)
errors=[]
QMessageBox.warning=lambda *args: errors.append(str(args[2]))
window.collection_service.client.max_retries=5
window.collection_service.client.request_interval=1.2
window.page_search.apply_filters(LicitacaoFilter(ufs=['GO','DF'],esferas=['M','D','E','F'],status_proposta='recebendo_proposta',prazo_dias=7))
results=[]
window.page_search._iniciar_coleta_online()
window.page_search.worker.collection_finished.connect(results.append)
wait(window.page_search.worker); wait(window.page_search.query_worker)
report={'tested_at':get_current_time_sp().isoformat(),'source':'PNCP oficial, sem fixtures','collection':results,'errors':errors}
repo=ProcurementRepository(db); all_items,total=repo.search(LicitacaoFilter(ufs=['GO','DF'],status_proposta='todas',limite=100000))
report['stored']=total; report['official_ids']=[x.numero_controle_pncp for x in all_items]
filters=LicitacaoFilter(ufs=['GO','DF'],status_proposta='recebendo_proposta',category_ids=[1],subcategory_ids=[1],prazo_dias=7)
window.page_search.apply_filters(filters); window.page_search._pesquisar_localmente(); wait(window.page_search.query_worker)
window.page_search.table_model.select_all(True); ids=window.page_search.table_model.get_selected_pncp_ids()
path=ExportService(repo).export_to_excel(DEST/'software_filtrados.xlsx',filters)
wb=load_workbook(path); report['software_count']=repo.search(filters)[1]; report['filtered_export_rows']=wb['LICITAÇÕES'].max_row-1; wb.close()
assert report['software_count']==report['filtered_export_rows']
selected=ExportService(repo).export_to_excel(DEST/'software_selecionados.xlsx',filters,ids)
wb=load_workbook(selected); report['selected_export_rows']=wb['LICITAÇÕES'].max_row-1; wb.close(); assert report['selected_export_rows']==len(ids)
# Arquivo completo real permite examinar categorias mesmo se a consulta específica for vazia.
ExportService(repo).export_to_excel(DEST/'todos_reais.xlsx',LicitacaoFilter(ufs=['GO','DF'],status_proposta='todas'))
window.grab().save(str(DEST/'pesquisa.png')); window._navigate(6); app.processEvents(); window.grab().save(str(DEST/'classificacao.png'))
window.close(); app.processEvents()
reopened=MainWindow(DatabaseConnection(db.db_path)); reopened.show(); wait(reopened.page_search.query_worker)
report['reopened_count']=reopened.procurement_repo.get_statistics()['total_registros']; assert report['reopened_count']==total
reopened.close(); app.processEvents()
with db.session() as c: report['integrity']=c.execute('PRAGMA integrity_check').fetchone()[0]; report['foreign_key_errors']=len(c.execute('PRAGMA foreign_key_check').fetchall())
(DEST/'resultado.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({k:v for k,v in report.items() if k!='official_ids'},ensure_ascii=False,indent=2))
if not results or results[0]['status']!='Concluída': sys.exit(2)
