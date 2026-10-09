import json
from datetime import datetime,timedelta
from unittest.mock import MagicMock
import httpx
import pytest
from openpyxl import load_workbook
from src.database.connection import DatabaseConnection
from src.database.migrations import init_database,SCHEMA_SQL
from src.database.repositories import ProcurementRepository
from src.domain.entities import Licitacao,Estado,Municipio
from src.domain.filters import LicitacaoFilter
from src.application.collection_service import CollectionService
from src.application.export_service import ExportService
from src.integrations.pncp_client import PNCPClient,PNCPClientError
from src.utils.datetime_utils import get_current_time_sp


def test_states_do_not_cascade_delete(location_repo):
    location_repo.save_estados([Estado('GO','Goiás',52)])
    location_repo.save_municipios([Municipio('5208707','Goiânia','GO')])
    location_repo.save_estados([Estado('GO','Goiás',52)])
    assert len(location_repo.get_municipios_por_uf('GO'))==1


def test_migration_preserves_existing(tmp_path):
    db=DatabaseConnection(tmp_path/'old.db')
    with db.session() as c:
        for sql in SCHEMA_SQL.split(';'):
            if sql.strip(): c.execute(sql)
    repo=ProcurementRepository(db); repo.upsert(Licitacao('old',objeto='oficial'))
    backup=db.backup()
    with db.session() as c: init_database(c)
    with db.session() as c: init_database(c); assert c.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
    assert repo.get_by_pncp_id('old').objeto=='oficial' and backup.exists()


def test_backup_restore_wal(db_conn,procurement_repo):
    procurement_repo.upsert_batch([Licitacao('one',objeto='notebooks')]); backup=db_conn.backup()
    procurement_repo.upsert_batch([Licitacao('two',objeto='medicamentos')])
    prior=db_conn.restore(backup)
    assert prior.exists() and procurement_repo.get_by_pncp_id('one')
    assert procurement_repo.get_by_pncp_id('two') is None


def test_atomic_rollback(db_conn):
    with pytest.raises(RuntimeError):
        with db_conn.session() as c:
            init_database(c)
            c.execute("INSERT INTO configuracoes VALUES ('should_rollback','yes')")
            raise RuntimeError()
    with db_conn.session() as c: assert c.execute("SELECT * FROM configuracoes WHERE chave='should_rollback'").fetchone() is None


def test_current_status_and_future_start(procurement_repo):
    now=get_current_time_sp()
    items=[Licitacao('open',data_encerramento_proposta=(now+timedelta(days=2)).isoformat(),situacao_compra_id=1),Licitacao('expired',data_encerramento_proposta=(now-timedelta(days=1)).isoformat()),Licitacao('future',data_abertura_proposta=(now+timedelta(days=1)).isoformat(),data_encerramento_proposta=(now+timedelta(days=2)).isoformat()),Licitacao('cancelled',situacao_compra_id=3,data_encerramento_proposta=(now+timedelta(days=2)).isoformat())]
    procurement_repo.upsert_batch(items)
    assert [x.numero_controle_pncp for x in procurement_repo.search(LicitacaoFilter())[0]]==['open']


@pytest.mark.parametrize('case',['repeat','empty','wrong_page','limit','invalid'])
def test_pagination_failures_not_success(case,procurement_repo,sync_repo,sample_raw_pncp):
    client=MagicMock(); p={'data':[sample_raw_pncp],'totalPaginas':2,'totalRegistros':2,'numeroPagina':1}
    second={'data':[],'totalPaginas':3,'numeroPagina':2}
    if case=='repeat': second={**p,'numeroPagina':2}
    if case=='wrong_page': second={**p,'numeroPagina':1}
    if case=='invalid': second={'message':'error'}
    client.get_propostas.side_effect=[p,second]
    result=CollectionService(procurement_repo,sync_repo,client).run_collection(LicitacaoFilter(ufs=['GO']),max_pages_per_partition=1 if case=='limit' else 500)
    assert result['status']=='Concluída com ressalvas' and result['erros']
    assert procurement_repo.get_statistics()['total_registros']==1


def test_publication_endpoint_and_local_filters(procurement_repo,sync_repo,sample_raw_pncp):
    client=MagicMock(); client.get_publicacoes.return_value={'data':[sample_raw_pncp],'totalPaginas':1,'totalRegistros':1}
    result=CollectionService(procurement_repo,sync_repo,client).run_collection(LicitacaoFilter(ufs=['GO'],esferas=['F'],status_proposta='todas',modalidades_ids=[6],data_publicacao_inicio='2026-10-01',data_publicacao_fim='2026-10-08'))
    assert result['status']=='Concluída' and result['novos']==0
    assert client.get_publicacoes.call_count==1 and not client.get_propostas.called


def test_http_204_and_parameters():
    with PNCPClient(request_interval=0) as client:
        client.client.close()
        def handler(request):
            assert 'dataInicial' not in request.url.params
            return httpx.Response(204)
        client.client=httpx.Client(transport=httpx.MockTransport(handler))
        assert client.get_propostas('20261031',data_inicial='20261001')['data']==[]


def test_excel_full_filter_types_metadata(procurement_repo,tmp_path):
    rows=[Licitacao(str(i),objeto='notebooks' if i%2 else 'medicamentos',valor_estimado=123.45,data_publicacao_pncp='2026-10-08T12:00:00',uf_sigla='GO') for i in range(12)]
    procurement_repo.upsert_batch(rows)
    f=LicitacaoFilter(status_proposta='todas',category_ids=[1],limite=2,orgao=None)
    dest=ExportService(procurement_repo).export_to_excel(tmp_path/'out.xlsx',f)
    wb=load_workbook(dest); ws=wb['LICITAÇÕES']
    assert ws.max_row==7 and isinstance(ws.cell(2,14).value,datetime)
    assert ws.cell(2,13).value==123.45 and ws.cell(2,24).value=='Tecnologia da Informação'
    empty=ExportService(procurement_repo).export_to_excel(tmp_path/'empty.xlsx',f,[])
    assert load_workbook(empty)['LICITAÇÕES'].max_row==1
    f.orgao='=1+1'
    dest=ExportService(procurement_repo).export_to_excel(tmp_path/'safe.xlsx',f)
    assert all(c.data_type!='f' for row in load_workbook(dest)['METADADOS'] for c in row)


def test_classification_sort(procurement_repo):
    procurement_repo.upsert_batch([Licitacao('hardware',objeto='notebooks'),Licitacao('medicine',objeto='medicamentos')])
    f=LicitacaoFilter(status_proposta='todas',ordenar_por='category_name')
    items,_=procurement_repo.search(f)
    assert [x.numero_controle_pncp for x in items]==['medicine','hardware']
    f.ordem_asc=False
    assert [x.numero_controle_pncp for x in procurement_repo.search(f)[0]]==['hardware','medicine']


def test_sync_updates_endpoint(procurement_repo,sync_repo,sample_raw_pncp):
    from src.application.synchronization_service import SynchronizationService
    from src.integrations.pncp_adapter import PNCPAdapter
    procurement_repo.upsert_batch([PNCPAdapter.parse_licitacao(sample_raw_pncp)])
    changed=dict(sample_raw_pncp,situacaoCompraId=3,situacaoCompraNome='Revogada')
    client=MagicMock(); client.get_atualizacoes.return_value={'data':[changed],'totalPaginas':1,'totalRegistros':1}
    f=LicitacaoFilter(ufs=['GO'],modalidades_ids=[6])
    result=SynchronizationService(procurement_repo,sync_repo,client).sync_open_proposals(f)
    assert result['status']=='Concluída' and result['atualizados']==1
    assert client.get_atualizacoes.called and not client.get_propostas.called
    assert procurement_repo.search(LicitacaoFilter())[1]==0


def test_bad_restore_preserves_data(db_conn,procurement_repo,tmp_path):
    procurement_repo.upsert(Licitacao('preserved',objeto='original'))
    bad=tmp_path/'bad.db'; bad.write_text('not sqlite')
    with pytest.raises(Exception): db_conn.restore(bad)
    assert procurement_repo.get_by_pncp_id('preserved').objeto=='original'


def test_cancel_during_backoff():
    client=PNCPClient(request_interval=0)
    with pytest.raises(PNCPClientError,match='cancelada'): client._wait(100,lambda:True)
    client.close()


def test_individual_official_endpoint():
    with PNCPClient(request_interval=0) as client:
        client.client.close()
        def handler(request):
            assert str(request.url)=='https://pncp.gov.br/api/consulta/v1/orgaos/00394411000109/compras/2026/9'
            return httpx.Response(200,json={'numeroControlePNCP':'official-test'})
        client.client=httpx.Client(transport=httpx.MockTransport(handler))
        assert client.get_compra('00394411000109',2026,9)['numeroControlePNCP']=='official-test'


def test_unclassified_origin_filter(procurement_repo):
    procurement_repo.upsert_batch([Licitacao('unknown',objeto='serviços diversos')])
    f=LicitacaoFilter(status_proposta='todas',classification_origin='automatic',only_unclassified=True)
    assert procurement_repo.search(f)[1]==1
    f.classification_origin='manual'
    assert procurement_repo.search(f)[1]==0

@pytest.mark.parametrize('text,expected',[('1.250,50',1250.5),('1250.50',1250.5),('1.000',1000),('',None),('0,00',0)])
def test_currency_input(text,expected):
    from src.domain.validation import parse_currency_input
    assert parse_currency_input(text)==expected

@pytest.mark.parametrize('text',['abc','NaN','1,2,3','-10','1.2.3'])
def test_invalid_currency_input(text):
    from src.domain.validation import parse_currency_input
    with pytest.raises(ValueError): parse_currency_input(text)
