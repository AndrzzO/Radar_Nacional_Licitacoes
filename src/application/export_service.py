"""Serviço de geração e exportação de planilhas Excel (.xlsx) profissionais via OpenPyXL."""
from pathlib import Path
from datetime import datetime
import os, tempfile
from src.classification.store import ClassificationStore
from src.classification.engine import ENGINE_VERSION
from typing import List, Optional, Callable
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.database.repositories import ProcurementRepository
from src.domain.entities import Licitacao
from src.domain.filters import LicitacaoFilter
from src.utils.datetime_utils import format_datetime_br, get_current_time_sp, parse_datetime, TZ_BRASILIA
from src.utils.file_utils import sanitize_for_excel
from src.utils.logging_config import get_logger

logger = get_logger()

# Cores e estilos visuais institucionais
HEADER_FILL = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")  # Azul marinho corporativo
HEADER_FONT = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Segoe UI", size=13, bold=True, color="1B365D")
META_LABEL_FONT = Font(name="Segoe UI", size=10, bold=True, color="333333")
BODY_FONT = Font(name="Segoe UI", size=9, color="111111")
LINK_FONT = Font(name="Segoe UI", size=9, color="004B87", underline="single")

BORDER_THIN = Border(
    left=Side(style="thin", color="E0E0E0"),
    right=Side(style="thin", color="E0E0E0"),
    top=Side(style="thin", color="E0E0E0"),
    bottom=Side(style="thin", color="E0E0E0"),
)

HEADER_COLUMNS = [
    ("ID PNCP", 35, Alignment(horizontal="left", vertical="center")),
    ("Número da Licitação", 20, Alignment(horizontal="left", vertical="center")),
    ("Ano", 10, Alignment(horizontal="center", vertical="center")),
    ("Modalidade", 22, Alignment(horizontal="left", vertical="center")),
    ("Esfera", 14, Alignment(horizontal="center", vertical="center")),
    ("Órgão", 35, Alignment(horizontal="left", vertical="center")),
    ("CNPJ", 18, Alignment(horizontal="center", vertical="center")),
    ("Unidade Compradora", 30, Alignment(horizontal="left", vertical="center")),
    ("UF", 8, Alignment(horizontal="center", vertical="center")),
    ("Município", 25, Alignment(horizontal="left", vertical="center")),
    ("Código IBGE", 14, Alignment(horizontal="center", vertical="center")),
    ("Objeto Completo", 50, Alignment(horizontal="left", vertical="center", wrap_text=True)),
    ("Valor Estimado (R$)", 20, Alignment(horizontal="right", vertical="center")),
    ("Data Publicação", 18, Alignment(horizontal="center", vertical="center")),
    ("Início das Propostas", 18, Alignment(horizontal="center", vertical="center")),
    ("Encerramento das Propostas", 20, Alignment(horizontal="center", vertical="center")),
    ("Situação Oficial", 22, Alignment(horizontal="left", vertical="center")),
    ("Situação do Prazo", 18, Alignment(horizontal="center", vertical="center")),
    ("Última Atualização Oficial", 20, Alignment(horizontal="center", vertical="center")),
    ("Link da Publicação", 35, Alignment(horizontal="left", vertical="center")),
    ("Link do Edital", 35, Alignment(horizontal="left", vertical="center")),
    ("Fonte de Dados", 14, Alignment(horizontal="center", vertical="center")),
    ("Data da Coleta Local", 20, Alignment(horizontal="center", vertical="center")),
    ("Categoria", 15, Alignment(horizontal="left", vertical="center")),
    ("Subcategoria", 15, Alignment(horizontal="left", vertical="center")),
]

for title in ('Categorias secundárias','Nível de confiança','Pontuação de classificação','Origem da classificação','Termos identificados','Data da classificação'):
    HEADER_COLUMNS.append((title, 25, Alignment(horizontal='left',vertical='center',wrap_text=True)))


def excel_date(value):
    dt=parse_datetime(value) if isinstance(value,str) else value
    return dt.astimezone(TZ_BRASILIA).replace(tzinfo=None) if dt else None


class ExportService:
    """Gerador profissional de arquivos Excel para licitações e metadados."""

    def __init__(self, procurement_repo: ProcurementRepository):
        self.repo = procurement_repo

    def export_to_excel(
        self,
        output_path: Path | str,
        filtros: LicitacaoFilter,
        selected_pncp_ids: Optional[List[str]] = None,
        on_progress: Optional[Callable[[int, int], None]] = None,
    ) -> Path:
        """
        Gera arquivo XLSX completo com as abas LICITAÇÕES e METADADOS.
        """
        dest_path = Path(output_path)
        dest_path.parent.mkdir(parents=True, exist_ok=True)

        logger.info(f"Iniciando exportação de registros para {dest_path}...")
        licitacoes = self.repo.get_all_for_export(filtros, selected_pncp_ids)
        total_items = len(licitacoes)

        wb = Workbook()

        # -------------------------------------------------------------
        # ABA 1: LICITAÇÕES
        # -------------------------------------------------------------
        ws_lic = wb.active
        ws_lic.title = "LICITAÇÕES"
        ws_lic.views.sheetView[0].showGridLines = True

        # Escrever Cabeçalhos
        for col_idx, (col_name, _, align) in enumerate(HEADER_COLUMNS, start=1):
            cell = ws_lic.cell(row=1, column=col_idx, value=col_name)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORDER_THIN
        ws_lic.row_dimensions[1].height = 28

        # Escrever Linhas de Dados
        for row_idx, lic in enumerate(licitacoes, start=2):
            if on_progress and (row_idx % 200 == 0 or row_idx == total_items + 1):
                on_progress(row_idx - 1, total_items)

            classification=getattr(lic,"classification",{})
            primary=classification.get("primary") or {}
            row_data = [
                sanitize_for_excel(lic.numero_controle_pncp),
                sanitize_for_excel(lic.numero_edital),
                lic.ano,
                sanitize_for_excel(lic.modalidade_nome),
                sanitize_for_excel(lic.esfera_nome),
                sanitize_for_excel(lic.orgao_nome),
                sanitize_for_excel(lic.orgao_cnpj),
                sanitize_for_excel(lic.unidade_nome),
                sanitize_for_excel(lic.uf_sigla),
                sanitize_for_excel(lic.municipio_nome),
                sanitize_for_excel(lic.municipio_ibge),
                sanitize_for_excel(lic.objeto),
                lic.valor_estimado,  # Float real ou None
                excel_date(lic.data_publicacao_pncp),
                excel_date(lic.data_abertura_proposta),
                excel_date(lic.data_encerramento_proposta),
                sanitize_for_excel(lic.situacao_compra_nome),
                sanitize_for_excel(lic.situacao_prazo),
                excel_date(lic.data_atualizacao_global or lic.data_atualizacao_pncp),
                sanitize_for_excel(lic.link_pncp),
                sanitize_for_excel(lic.link_edital),
                sanitize_for_excel(lic.fonte_dados),
                excel_date(lic.data_inclusao_local),
                sanitize_for_excel(primary.get('category','Não Classificado')),
                sanitize_for_excel(primary.get('subcategory','')),
                sanitize_for_excel('; '.join(c['category'] for c in classification.get('classifications',[]) if not c['primary'])),
                primary.get('confidence','Não Classificado'),
                primary.get('score',0),
                classification.get('origin','Pendente'),
                sanitize_for_excel('; '.join(e['matched'] for c in classification.get('classifications',[]) for e in c.get('evidence',[]))),
                excel_date(classification.get('date')),
            ]

            for col_idx, (val, (_, _, align)) in enumerate(zip(row_data, HEADER_COLUMNS), start=1):
                cell = ws_lic.cell(row=row_idx, column=col_idx, value=val)
                cell.font = BODY_FONT
                cell.alignment = align
                cell.border = BORDER_THIN

                if isinstance(val,datetime):
                    cell.number_format = "dd/mm/yyyy hh:mm"

                # Formatação de número monetário
                if col_idx == 13 and isinstance(val, (int, float)):
                    cell.number_format = 'R$ #,##0.00'

                # Formatação de hiperlinks clicáveis
                if col_idx in (20, 21) and val and str(val).startswith("http"):
                    cell.hyperlink = str(val)
                    cell.font = LINK_FONT

            ws_lic.row_dimensions[row_idx].height = 20

        # Ativar AutoFiltro e Congelar Cabeçalho
        last_col_letter = get_column_letter(len(HEADER_COLUMNS))
        last_row = max(len(licitacoes) + 1, 1)
        ws_lic.auto_filter.ref = f"A1:{last_col_letter}{last_row}"
        ws_lic.freeze_panes = "A2"

        # Ajustar Largura das Colunas
        for col_idx, (_, col_width, _) in enumerate(HEADER_COLUMNS, start=1):
            col_letter = get_column_letter(col_idx)
            ws_lic.column_dimensions[col_letter].width = col_width

        # -------------------------------------------------------------
        # ABA 2: METADADOS
        # -------------------------------------------------------------
        ws_meta = wb.create_sheet(title="METADADOS")
        ws_meta.views.sheetView[0].showGridLines = True

        stats = self.repo.get_statistics()
        agora_sp = get_current_time_sp()

        catalog,version,rules_date=ClassificationStore(self.repo.db).catalog()
        metadata_rows = [
            ("RADAR NACIONAL DE LICITAÇÕES — RELATÓRIO DE EXPORTAÇÃO", ""),
            ("", ""),
            ("Data e Hora da Exportação", excel_date(agora_sp)),
            ("Fonte Oficial de Dados", "Portal Nacional de Contratações Públicas (PNCP) - API Oficial"),
            ("Quantidade de Registros Exportados", total_items),
            ("Tipo de Seleção", "Registros selecionados manualmente" if selected_pncp_ids is not None else "Todos os registros filtrados"),
            ("Estados (UFs) Filtrados", ", ".join(filtros.ufs) if filtros.ufs else "Brasil Inteiro (Todas as UFs)"),
            ("Municípios (IBGE) Filtrados", ", ".join(filtros.municipios_ibge) if filtros.municipios_ibge else "Todos os municípios"),
            ("Esferas Administrativas", ", ".join(filtros.esferas) if filtros.esferas else "Todas as esferas"),
            ("Situação do Prazo / Propostas", filtros.status_proposta),
            ("Período de Publicação", f"{filtros.data_publicacao_inicio or 'Início'} até {filtros.data_publicacao_fim or 'Atual'}"),
            ("Período de Encerramento de Propostas", f"{filtros.data_encerramento_inicio or 'Início'} até {filtros.data_encerramento_fim or 'Atual'}"),
            ("Filtro por Palavra-chave", filtros.palavra_chave or "Nenhum"),
            ("Filtro por Órgão Público", filtros.orgao or "Nenhum"),
            ("Data da Última Sincronização do Banco", excel_date(stats.get("data_ultima_sinc")) or "Não realizada"),
            ("", ""),
            ("OBSERVAÇÕES TÉCNICAS E INTEGRIDADE:", ""),
            ("1. Os dados foram extraídos diretamente dos serviços oficiais documentados do PNCP.", ""),
            ("2. Valores ausentes na fonte não foram substituídos por estimativas inventadas ou R$ 0,00.", ""),
            ("3. Textos foram sanitizados contra injeção de fórmulas arbitrárias em conformidade com normas de segurança.", ""),
            ("4. A categorização é realizada automaticamente por regras de palavras-chave e pode apresentar divergências. A descrição e as condições oficiais devem ser confirmadas no edital.", ""),
            ("Modalidades", ', '.join(map(str,filtros.modalidades_ids)) or 'Todas'),
            ("Categorias", ', '.join(c['name'] for c in catalog['categories'] if c['id'] in filtros.category_ids) or 'Todas'),
            ("Subcategorias", ', '.join(c['name'] for c in catalog['subcategories'] if c['id'] in filtros.subcategory_ids) or 'Todas'),
            ("Níveis de confiança", ', '.join(filtros.confidence_levels) or 'Todos'),
            ("Versão do classificador", ENGINE_VERSION),
            ("Versão do catálogo atual", version),
            ("Versões presentes nos registros", ', '.join(map(str,sorted({getattr(x,'classification',{}).get('version',0) for x in licitacoes})))),
            ("Atualização das regras", excel_date(rules_date)),
            ("Filtros completos (JSON)", filtros.to_json()),
        ]

        for r_idx, (k, v) in enumerate(metadata_rows, start=1):
            cell_k = ws_meta.cell(row=r_idx, column=1, value=sanitize_for_excel(k))
            cell_v = ws_meta.cell(row=r_idx, column=2, value=v if isinstance(v,datetime) else sanitize_for_excel(v))

            if isinstance(v,datetime): cell_v.number_format="dd/mm/yyyy hh:mm"
            if r_idx == 1:
                cell_k.font = TITLE_FONT
            elif "OBSERVAÇÕES" in k:
                cell_k.font = TITLE_FONT
            elif k.startswith(("1.", "2.", "3.", "4.")):
                cell_k.font = BODY_FONT
            else:
                cell_k.font = META_LABEL_FONT
                cell_v.font = BODY_FONT

        ws_meta.column_dimensions["A"].width = 38
        ws_meta.column_dimensions["B"].width = 65

        # Salvar arquivo
        fd,temp_path=tempfile.mkstemp(suffix='.xlsx',dir=dest_path.parent)
        os.close(fd)
        try:
            wb.save(temp_path)
            os.replace(temp_path,dest_path)
        finally:
            wb.close()
            if os.path.exists(temp_path): os.unlink(temp_path)
        logger.info(f"Exportação para {dest_path} concluída com sucesso ({total_items} linhas).")
        return dest_path
